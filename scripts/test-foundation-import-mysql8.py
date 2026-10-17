#!/usr/bin/env python3
"""Run the real foundation import helpers against disposable local MySQL 8 only."""
import argparse
import json
import os
from pathlib import Path
import shutil
import subprocess
import tempfile
import time

project = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser()
parser.add_argument(
    "--binaries-root",
    default="/home/ubuntu/foundation-producao-20261016/mysql8-test-binaries/root",
)
parser.add_argument(
    "--output",
    default="/home/ubuntu/foundation-excel-v2-20261017/mysql-integration",
)
args = parser.parse_args()

root = Path(args.binaries_root).resolve()
mysql = root / "usr/bin/mysql"
mysqld = root / "usr/sbin/mysqld"
if not mysql.is_file() or not mysqld.is_file():
    raise SystemExit("MYSQL8_DISPOSABLE_BINARIES_REQUIRED")

out = Path(args.output).resolve()
out.mkdir(parents=True, exist_ok=True)
work = Path(tempfile.mkdtemp(prefix="idefazei-import-v2-mysql8-", dir="/tmp"))
work.chmod(0o700)
data = work / "data"
data.mkdir(mode=0o700)
socket = work / "mysql.sock"

env = os.environ.copy()
env["LD_LIBRARY_PATH"] = str(root / "usr/lib/x86_64-linux-gnu")
env.pop("DATABASE_URL", None)
env["FOUNDATION_IMPORT_DISPOSABLE_MYSQL"] = "sandbox-only"
env["FOUNDATION_DISPOSABLE_SOCKET"] = str(socket)
env["CI"] = "1"

server = None
exit_code = 1
report = {
    "kind": "disposable-real-mysql-foundation-import-v2",
    "socket_only": True,
    "network_disabled": True,
    "vps_access": False,
    "database": "foundation_import_v2_test",
    "test_exit_code": None,
}


def query(sql, database=None):
    if database is not None and database != "foundation_import_v2_test":
        raise RuntimeError("DISPOSABLE_DATABASE_REQUIRED")
    command = [
        str(mysql),
        "--no-defaults",
        "--socket=" + str(socket),
        "--user=root",
        "--batch",
        "--skip-column-names",
    ]
    if database:
        command += ["--database=" + database]
    result = subprocess.run(
        command,
        input=sql,
        text=True,
        capture_output=True,
        env=env,
        timeout=30,
    )
    if result.returncode:
        raise RuntimeError("DISPOSABLE_MYSQL_QUERY_FAILED")
    return result.stdout.strip()


schema = r"""
CREATE TABLE courses (
 id INT PRIMARY KEY AUTO_INCREMENT,
 churchId INT NOT NULL,
 name VARCHAR(255) NOT NULL,
 type ENUM('salvacao','oracao','biblia','igreja','espirito_santo','batismo','outro') NOT NULL,
 description TEXT NULL,
 active BOOLEAN NOT NULL DEFAULT 1,
 createdAt TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP
) ENGINE=InnoDB;
CREATE TABLE library_items (
 id INT PRIMARY KEY AUTO_INCREMENT,
 churchId INT NOT NULL,
 title VARCHAR(255) NOT NULL,
 type ENUM('pdf','video','apostila','devocional') NOT NULL,
 fileUrl TEXT NULL,
 thumbnailUrl TEXT NULL,
 description TEXT NULL,
 requiredRole ENUM('pastor_presidente','pastor_local','supervisor','lider','consolidador','diacono','secretario','tesoureiro','membro') NOT NULL DEFAULT 'membro',
 createdAt TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP
) ENGINE=InnoDB;
CREATE TABLE foundation_modules (
 id INT PRIMARY KEY AUTO_INCREMENT,
 churchId INT NOT NULL,
 courseId INT NOT NULL,
 title VARCHAR(160) NOT NULL,
 description VARCHAR(500) NULL,
 position INT NOT NULL DEFAULT 0,
 active BOOLEAN NOT NULL DEFAULT 1,
 createdByChurchUserId INT NOT NULL,
 createdAt TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
 updatedAt TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
 INDEX foundation_modules_church_course_position_idx(churchId,courseId,position)
) ENGINE=InnoDB;
CREATE TABLE foundation_studies (
 id INT PRIMARY KEY AUTO_INCREMENT,
 churchId INT NOT NULL,
 courseId INT NOT NULL,
 moduleId INT NULL,
 title VARCHAR(160) NOT NULL,
 weekStart DATE NULL,
 summary VARCHAR(500) NULL,
 content TEXT NULL,
 position INT NOT NULL DEFAULT 0,
 active BOOLEAN NOT NULL DEFAULT 1,
 createdByChurchUserId INT NOT NULL,
 clientRequestId VARCHAR(36) NULL,
 requestHash CHAR(64) NULL,
 requestSchemaVersion INT NULL,
 createdAt TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
 updatedAt TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
 INDEX foundation_studies_church_course_position_idx(churchId,courseId,position),
 INDEX foundation_studies_church_module_position_idx(churchId,moduleId,position)
) ENGINE=InnoDB;
CREATE TABLE foundation_study_blocks (
 id INT PRIMARY KEY AUTO_INCREMENT,
 churchId INT NOT NULL,
 studyId INT NOT NULL,
 title VARCHAR(160) NOT NULL,
 content TEXT NOT NULL,
 reflectionPrompt TEXT NULL,
 position INT NOT NULL DEFAULT 0,
 active BOOLEAN NOT NULL DEFAULT 1,
 createdByChurchUserId INT NOT NULL,
 createdAt TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
 updatedAt TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
 INDEX foundation_study_blocks_church_study_position_idx(churchId,studyId,position)
) ENGINE=InnoDB;
CREATE TABLE foundation_questions (
 id INT PRIMARY KEY AUTO_INCREMENT,
 churchId INT NOT NULL,
 blockId INT NOT NULL,
 prompt TEXT NOT NULL,
 options JSON NOT NULL,
 correctOptionId VARCHAR(80) NOT NULL,
 explanation TEXT NOT NULL,
 position INT NOT NULL DEFAULT 0,
 active BOOLEAN NOT NULL DEFAULT 1,
 createdByChurchUserId INT NOT NULL,
 clientRequestId VARCHAR(36) NULL,
 requestHash CHAR(64) NULL,
 requestSchemaVersion INT NULL,
 createdAt TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
 updatedAt TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
 INDEX foundation_questions_church_block_position_idx(churchId,blockId,position)
) ENGINE=InnoDB;
CREATE TABLE foundation_study_materials (
 id INT PRIMARY KEY AUTO_INCREMENT,
 churchId INT NOT NULL,
 studyId INT NOT NULL,
 libraryItemId INT NOT NULL,
 position INT NOT NULL DEFAULT 0,
 notes TEXT NULL,
 createdAt TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
 INDEX foundation_study_materials_church_study_position_idx(churchId,studyId,position)
) ENGINE=InnoDB;
CREATE TABLE foundation_classes (
 id INT PRIMARY KEY AUTO_INCREMENT,
 churchId INT NOT NULL,
 courseId INT NOT NULL,
 studyId INT NOT NULL,
 classDate DATE NOT NULL,
 status ENUM('planejada','realizada','cancelada') NOT NULL DEFAULT 'planejada',
 notes TEXT NULL,
 createdByChurchUserId INT NOT NULL,
 updatedByChurchUserId INT NOT NULL,
 createdAt TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
 updatedAt TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
 INDEX foundation_classes_church_study_date_idx(churchId,studyId,classDate)
) ENGINE=InnoDB;
CREATE TABLE foundation_import_drafts (
 id INT PRIMARY KEY AUTO_INCREMENT,
 churchId INT NOT NULL,
 courseId INT NOT NULL,
 createdByChurchUserId INT NOT NULL,
 sourceFilename VARCHAR(255) NOT NULL,
 fileSha256 VARCHAR(64) NOT NULL,
 status ENUM('pendente','confirmado','cancelado','expirado') NOT NULL DEFAULT 'pendente',
 payload JSON NOT NULL,
 summary JSON NOT NULL,
 expiresAt TIMESTAMP NOT NULL,
 createdAt TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
 confirmedAt TIMESTAMP NULL,
 INDEX foundation_import_drafts_church_user_status_idx(churchId,createdByChurchUserId,status),
 INDEX foundation_import_drafts_church_expires_idx(churchId,expiresAt)
) ENGINE=InnoDB;
-- Sentinel tables: confirmation must not create learner answers, progress, or attendance.
CREATE TABLE foundation_question_attempts (
 id INT PRIMARY KEY AUTO_INCREMENT,
 churchId INT NOT NULL,
 enrollmentId INT NOT NULL,
 questionId INT NOT NULL,
 selectedOptionId VARCHAR(80) NOT NULL,
 isCorrect BOOLEAN NOT NULL,
 attemptNumber INT NOT NULL,
 clientAttemptId VARCHAR(80) NULL,
 answeredAt TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP
) ENGINE=InnoDB;
CREATE TABLE foundation_lesson_progress (
 id INT PRIMARY KEY AUTO_INCREMENT,
 churchId INT NOT NULL,
 enrollmentId INT NOT NULL,
 studyId INT NOT NULL,
 status ENUM('nao_iniciada','em_andamento','concluida') NOT NULL DEFAULT 'nao_iniciada',
 lastBlockPosition INT NOT NULL DEFAULT 0,
 reflection TEXT NULL,
 createdAt TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
 updatedAt TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP
) ENGINE=InnoDB;
CREATE TABLE foundation_block_progress (
 id INT PRIMARY KEY AUTO_INCREMENT,
 churchId INT NOT NULL,
 enrollmentId INT NOT NULL,
 studyId INT NOT NULL,
 blockId INT NOT NULL,
 status ENUM('em_andamento','concluida') NOT NULL DEFAULT 'em_andamento',
 reflection TEXT NULL,
 createdAt TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
 updatedAt TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP
) ENGINE=InnoDB;
CREATE TABLE foundation_class_attendance (
 id INT PRIMARY KEY AUTO_INCREMENT,
 churchId INT NOT NULL,
 classId INT NOT NULL,
 enrollmentId INT NOT NULL,
 status ENUM('presente','ausente','justificado') NOT NULL,
 notes TEXT NULL,
 recordedByChurchUserId INT NOT NULL,
 createdAt TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP
) ENGINE=InnoDB;
"""

try:
    from openpyxl import load_workbook
    workbook = load_workbook(project / "client/src/documents/foundation/gabarito-escola-fundamentos-v2-vazio.xlsx")
    workbook["ESTUDO"]["B5"] = "Conhecendo o amor de Deus — importação QA"
    workbook["ESTUDO"]["C5"] = "19/10/2026"
    for row in range(5, 9):
        workbook["BLOCOS"].cell(row, 3, f"Conteúdo sintético da etapa {row - 4}.\n\nSegundo parágrafo preservado.")
    workbook["BLOCOS"]["D5"] = "Como esse ensino orienta suas atitudes?"
    values = [1, "Qual atitude corresponde ao ensino?", "Ouvir com atenção", "Ignorar as pessoas", None, None, None, None, "A", "Ouvir é uma atitude de cuidado."]
    for column, value in enumerate(values, 1):
        workbook["PERGUNTAS"].cell(5, column).value = value
    filled = work / "modelo-preenchido-sintetico.xlsx"
    workbook.save(filled)
    env["FOUNDATION_IMPORT_FILLED_FIXTURE"] = str(filled)
    report["official_workbook_filled_locally"] = True
    print("IMPORT_MYSQL8_INITIALIZING", flush=True)
    initialized = subprocess.run(
        [
            str(mysqld),
            "--no-defaults",
            "--initialize-insecure",
            "--basedir=" + str(root / "usr"),
            "--datadir=" + str(data),
            "--log-error=" + str(work / "init.log"),
        ],
        env=env,
        capture_output=True,
        timeout=45,
    )
    if initialized.returncode:
        raise RuntimeError("DISPOSABLE_MYSQL_INIT_FAILED")

    server = subprocess.Popen(
        [
            str(mysqld),
            "--no-defaults",
            "--basedir=" + str(root / "usr"),
            "--datadir=" + str(data),
            "--socket=" + str(socket),
            "--pid-file=" + str(work / "mysql.pid"),
            "--skip-networking",
            "--mysqlx=OFF",
            "--secure-file-priv=" + str(work),
            "--log-error=" + str(work / "server.log"),
            "--innodb-buffer-pool-size=64M",
        ],
        env=env,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    deadline = time.monotonic() + 25
    while True:
        if server.poll() is not None:
            raise RuntimeError("DISPOSABLE_MYSQL_START_FAILED")
        try:
            version = query("SELECT VERSION();")
            break
        except RuntimeError:
            if time.monotonic() > deadline:
                raise
            time.sleep(0.3)
    if not version.startswith("8.0."):
        raise RuntimeError("MYSQL8_REQUIRED")
    if query("SELECT @@socket;") != str(socket):
        raise RuntimeError("DISPOSABLE_SOCKET_MISMATCH")
    query("CREATE DATABASE foundation_import_v2_test;")
    query(schema, "foundation_import_v2_test")
    report["mysql_version"] = version
    report["socket"] = str(socket)
    print("IMPORT_MYSQL8_READY_LOCAL_ONLY=" + version, flush=True)

    command = [
        "pnpm",
        "exec",
        "vitest",
        "run",
        "server/foundation-import-mysql.integration.test.ts",
        "--reporter=verbose",
        "--reporter=json",
        "--outputFile.json=" + str(out / "vitest.json"),
    ]
    with (out / "vitest.log").open("w") as logfile:
        result = subprocess.run(
            command,
            cwd=project,
            env=env,
            stdout=logfile,
            stderr=subprocess.STDOUT,
            timeout=240,
        )
    exit_code = result.returncode
    report["test_exit_code"] = exit_code
    if (out / "vitest.json").is_file():
        totals = json.loads((out / "vitest.json").read_text())
        report.update(
            {
                "tests_pass": totals.get("numPassedTests"),
                "tests_fail": totals.get("numFailedTests"),
                "tests_pending": totals.get("numPendingTests"),
                "test_files_pass": totals.get("numPassedTestSuites"),
                "test_files_fail": totals.get("numFailedTestSuites"),
            }
        )
except Exception as error:
    report["runner_error"] = str(error) if isinstance(error, (RuntimeError, AssertionError)) else type(error).__name__
    if (work / "server.log").is_file():
        shutil.copyfile(work / "server.log", out / "disposable-server-diagnostic.log")
finally:
    if server is not None and server.poll() is None:
        try:
            query("SHUTDOWN;")
        except Exception:
            server.terminate()
        try:
            server.wait(timeout=15)
        except subprocess.TimeoutExpired:
            server.kill()
            server.wait(timeout=5)
    shutil.rmtree(work, ignore_errors=True)
    report["process_stopped"] = server is None or server.poll() is not None
    report["temporaries_removed"] = not work.exists()
    (out / "summary.json").write_text(json.dumps(report, indent=2) + "\n")
    print(json.dumps(report, indent=2), flush=True)
    print("IMPORT_MYSQL8_STOPPED_AND_REMOVED", flush=True)

raise SystemExit(exit_code)
